#!/usr/bin/env python3
"""Export the manuscript's rendered HTML tables as editable Word tables.

This is a document conversion only. It copies displayed strings and images;
it never reads research data, evaluates models, or recalculates statistics.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import html
from PIL import Image, ImageColor
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LABELS = [
    ("Table_1", "Participant and measurement characteristics"),
    ("Table_2", "Recommendation adherence"),
    ("Table_3", "Near-eye metrics and geographic and photoperiod context"),
    ("Table_S1", "Descriptive sample flow"),
    ("Table_S2", "Full near-eye metric dictionary and distributions"),
    ("Table_S3", "Descriptive recommendation-range fractions"),
    ("Table_S4", "Exploratory cross-window recommendation-adherence associations"),
    ("Table_S5", "Near-eye fitted-curve dispersion and full-model R² allocation"),
    ("Table_S6", "Chest fitted-curve dispersion and full-model R² allocation"),
    ("Table_S7", "Geographic and photoperiod associations"),
    ("Table_S8", "Nonlinear civil-photoperiod classifications"),
    ("Table_S9", "Participant-hour routine associations"),
    ("Table_S10", "Person-level evidence synthesis"),
    ("Table_S11a", "Light-exposure behaviour and awareness, part 1"),
    ("Table_S11b", "Light-exposure behaviour and awareness, part 2"),
    ("Table_S12", "Visual light sensitivity"),
    ("Table_S13", "Chronotype and timing"),
    ("Table_S14", "Age and biological sex"),
    ("Table_S15", "Biological-sex-specific daily patterns and global tests"),
]

# Relative widths are layout choices, not transformations of table values.
WIDTHS = {
    "Table_1": [2.10] + [0.85] * 10,
    "Table_2": [1.00, 1.75, 1.80, 1.55, 1.55, 1.55, 0.90],
    "Table_3": [130, 190, 140, 110, 150, 235, 265],
    "Table_S1": [1.75, 3.65, 1.35],
    "Table_S2": [200, 60, 103] + [96] * 9 + [78, 185],
    "Table_S3": [1.35] + [1.15] * 8,
    "Table_S4": [1.0, 1.60, 0.80, 2.35, 4.00],
    "Table_S5": [2.45, 2.50, 2.85, 2.70],
    "Table_S6": [2.45, 2.50, 2.85, 2.70],
    "Table_S7": [1.95, 0.80, 1.00, 0.80, 1.00, 0.80, 0.90, 2.00],
    "Table_S8": [2.05, 0.70, 1.20, 2.00, 3.80],
    "Table_S9": [3.50, 2.00, 1.20],
    "Table_S10": [1.15, 1.40, 2.20, 1.30, 1.20, 1.45, 2.30],
    "Table_S11a": [2.20, 2.00, 2.00, 2.00, 2.00],
    "Table_S11b": [2.20, 2.00, 2.00, 2.00, 2.00],
    "Table_S12": [2.00, 1.35, 2.00, 1.05, 2.00, 0.65],
    "Table_S13": [2.40, 0.85, 2.00, 1.05, 3.10],
    "Table_S14": [1.50, 1.75, 1.85, 2.20, 1.00, 2.30],
    "Table_S15": [0.90, 1.20, 2.30, 0.80, 0.90, 1.00, 0.70, 0.95, 1.00],
}


def normalise(text):
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def compact(text):
    return re.sub(r"\s+", "", text.replace("\u00a0", " "))


def css(element):
    return dict(re.findall(r"([\w-]+)\s*:\s*([^;]+)", element.get("style", "")))


def visually_hidden(element):
    styles = css(element)
    classes = element.get("class", "").split()
    return (styles.get("display") == "none" or "clip-path" in styles
            or "sr-only" in classes or "visually-hidden" in classes)


def visible_text(element):
    copy = deepcopy(element)
    for node in list(copy.iterdescendants()):
        if isinstance(node.tag, str) and visually_hidden(node):
            node.drop_tree()
    return "".join(copy.itertext())


def colour(value, base="FFFFFF"):
    if not value or value in ("transparent", "inherit", "initial"):
        return None
    try:
        if value.startswith("rgba"):
            r, g, b, a = map(float, re.findall(r"[\d.]+", value))
            background = ImageColor.getrgb("#" + base)
            return "".join(f"{round(x*a+y*(1-a)):02X}" for x, y in zip((r,g,b), background))
        return "".join(f"{x:02X}" for x in ImageColor.getrgb(value)[:3])
    except (ValueError, TypeError):
        return None


def element_property(parent, tag, attributes):
    item = parent.find(qn(tag))
    if item is None:
        item = OxmlElement(tag)
        parent.append(item)
    for name, value in attributes.items():
        item.set(qn(name), str(value))
    return item


def format_paragraph(paragraph, align="left", keep=False):
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.keep_with_next = keep
    pf.keep_together = True
    pf.first_line_indent = Inches(0)


def add_run(paragraph, text, style):
    if not text:
        return
    text = re.sub(r"\s+", " ", text)
    if style.get("nowrap"):
        text = text.replace(" ", "\u00a0")
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(style.get("size", 10))
    run.bold = style.get("bold", False)
    run.italic = style.get("italic", False)
    run.font.superscript = style.get("sup", False)
    run.font.subscript = style.get("sub", False)
    if style.get("colour"):
        run.font.color.rgb = RGBColor.from_string(style["colour"])
    return run


def render_inline(element, paragraph, style, width, source_directory):
    if visually_hidden(element):
        return
    current = dict(style)
    styles = css(element)
    if (current.get("s2_numeric") and styles.get("white-space") == "nowrap"
            and "±" in "".join(element.itertext())):
        current["nowrap"] = True
    if styles.get("font-weight") in ("bold", "600", "700", "800") or element.tag in ("b", "strong"):
        current["bold"] = True
    if styles.get("font-style") == "italic" or element.tag in ("em", "i"):
        current["italic"] = True
    if colour(styles.get("color")):
        current["colour"] = colour(styles["color"])
    if element.tag == "sup":
        current.update(sup=True, size=8)
    if element.tag == "sub":
        current.update(sub=True, size=8)
    # The source's compact explanatory text remains subordinate, but legible.
    if "font-size" in styles and ("smaller" in styles["font-size"] or styles["font-size"] in ("9px", "10px", "11px", "85%", "0.85em", "0.8em")):
        current["size"] = 9
    if element.tag == "br":
        paragraph.add_run().add_break()
        return
    if element.tag == "img":
        src = element.get("src", "")
        if src.startswith("data:"):
            raw = base64.b64decode(src.split(",", 1)[1])
        else:
            raw = (source_directory / src).read_bytes()
        with Image.open(io.BytesIO(raw)) as img:
            image_width, image_height = img.size
        available = min(width - 0.10, 1.45)
        run = paragraph.add_run()
        shape = run.add_picture(io.BytesIO(raw), width=Inches(available), height=Inches(available*image_height/image_width))
        preceding = element.getprevious()
        description = (normalise("".join(preceding.itertext()))
                       if preceding is not None and visually_hidden(preceding) else None)
        shape._inline.docPr.set("descr", element.get("alt") or element.get("aria-label")
                               or description or "Distribution shown in the source table")
        return
    add_run(paragraph, element.text, current)
    for child in element:
        block = child.tag in ("p", "div")
        if block and normalise(paragraph.text):
            paragraph.add_run().add_break()
        render_inline(child, paragraph, current, width, source_directory)
        add_run(paragraph, child.tail, current)


def source_caption(table, index):
    if index < 3:
        ancestors = table.xpath("ancestor::div[contains(concat(' ',@class,' '),' quarto-float ')]")
        if ancestors:
            captions = ancestors[-1].xpath(".//*[contains(concat(' ',@class,' '),' quarto-float-caption ')]")
            if captions:
                return captions[-1]
        captions = table.xpath("ancestor::figure[1]/figcaption")
        return captions[-1] if captions else None
    containers = table.xpath("ancestor::div[starts-with(@id,'supp-table-s')]")
    paragraphs = containers[-1].xpath("./p") if containers else []
    return paragraphs[-1] if paragraphs else None


def grid_records(table):
    rows = table.xpath("./thead/tr|./tbody/tr")
    ncols = max(sum(int(cell.get("colspan", 1)) for cell in row.xpath("./th|./td")) for row in rows)
    occupied = set()
    records = []
    for r, row in enumerate(rows):
        column = 0
        for cell in row.xpath("./th|./td"):
            while (r, column) in occupied:
                column += 1
            colspan, rowspan = int(cell.get("colspan", 1)), int(cell.get("rowspan", 1))
            if column + colspan > ncols:
                raise RuntimeError("Invalid source table grid")
            records.append((r, column, rowspan, colspan, cell))
            for rr in range(r, r+rowspan):
                for cc in range(column, column+colspan):
                    occupied.add((rr, cc))
            column += colspan
    return rows, ncols, records


def export_one(source_table, source_path, reference, outdir, index, s2_numeric_reduction_pt=0):
    key, label = LABELS[index]
    rows, ncols, records = grid_records(source_table)
    doc = Document(reference)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)
    section = doc.sections[0]
    wide = key not in ("Table_S1", "Table_S9")
    section.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
    if key == "Table_S2":
        section.page_width, section.page_height = Inches(16.54), Inches(11.69)
    elif wide:
        section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    else:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.55)
    section.top_margin = section.bottom_margin = Inches(0.55)
    for ln in doc.element.xpath(".//w:lnNumType"):
        ln.getparent().remove(ln)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.line_spacing = 1.0
    title = doc.add_paragraph()
    format_paragraph(title, keep=True)
    title.paragraph_format.space_after = Pt(8)
    add_run(title, key.replace("_", " ") + ". " + label, {"size":14,"bold":True,"colour":"000000"})
    available = (section.page_width-section.left_margin-section.right_margin)/914400
    weights = WIDTHS[key]
    if len(weights) != ncols:
        raise RuntimeError(f"Width-grid mismatch in {key}: {len(weights)} versus {ncols}")
    widths = [available*w/sum(weights) for w in weights]
    word_table = doc.add_table(rows=len(rows), cols=ncols)
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    word_table.autofit = False
    props = word_table._tbl.tblPr
    element_property(props, "w:tblW", {"w:w":round(available*1440),"w:type":"dxa"})
    borders = element_property(props,"w:tblBorders",{})
    for edge in ("top","bottom","left","right","insideH","insideV"):
        element_property(borders,"w:"+edge,{"w:val":"single","w:sz":3,"w:color":"D9D9D9"})
    margins = element_property(props,"w:tblCellMar",{})
    for edge, size in (("top",65),("bottom",65),("left",55),("right",55)):
        element_property(margins,"w:"+edge,{"w:w":size,"w:type":"dxa"})
    for column, width in zip(word_table.columns,widths):
        column.width=Inches(width)
    for row in word_table.rows:
        element_property(row._tr.get_or_add_trPr(),"w:cantSplit",{})
        for c,width in zip(row.cells,widths):
            c.width=Inches(width)
    originals=[]
    s2_nowrap_spans_verified=0
    destination_cells=[]
    for r,c,rs,cs,source in records:
        cell=word_table.cell(r,c)
        if rs>1 or cs>1:
            cell=cell.merge(word_table.cell(r+rs-1,c+cs-1))
        cell.text=""
        width=sum(widths[c:c+cs])
        cell.width=Inches(width)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        heading=rows[r].getparent().tag=="thead"
        group="gt_group_heading" in source.get("class","")
        if heading:
            element_property(word_table.rows[r]._tr.get_or_add_trPr(),"w:tblHeader",{})
        styles=css(source)
        shade="F1F3F5" if heading or group else ("F6F6F6" if r%2==0 else "FFFFFF")
        bg=colour(styles.get("background-color"),shade)
        if bg:
            shade=bg
        element_property(cell._tc.get_or_add_tcPr(),"w:shd",{"w:fill":shade,"w:val":"clear"})
        align=styles.get("text-align")
        if not align:
            classes=source.get("class","")
            align="center" if "gt_center" in classes else "right" if "gt_right" in classes else "left"
        if group:
            align="left"
        p=cell.paragraphs[0]
        format_paragraph(p,align,keep=group)
        style={"size":10,"bold":heading or group,"colour":colour(styles.get("color")) or "333333"}
        if key=="Table_S2" and not heading and not group and cs==1 and 2<=c<=11:
            style.update(size=10-s2_numeric_reduction_pt,s2_numeric=True)
        render_inline(source,p,style,width,source_path.parent)
        if style.get("s2_numeric"):
            spans=[node for node in source.iterdescendants()
                   if node.tag=="span" and css(node).get("white-space")=="nowrap"
                   and "±" in "".join(node.itertext())]
            for span in spans:
                token=normalise("".join(span.itertext())).replace(" ","\u00a0")
                if token not in p.text:
                    raise RuntimeError("Complete S2 mean-plus-minus-SD unit lost its no-break spaces")
            s2_nowrap_spans_verified+=len(spans)
        originals.append(compact(visible_text(source)))
        destination_cells.append(cell)
    notes_and_caption=[]
    for row in source_table.xpath("./tfoot/tr"):
        for source in row.xpath("./th|./td"):
            p=doc.add_paragraph()
            format_paragraph(p,keep=True)
            p.paragraph_format.space_before=Pt(4)
            render_inline(source,p,{"size":9,"colour":"444444"},available,source_path.parent)
            notes_and_caption.append((source,p))
    caption=source_caption(source_table,index)
    if caption is not None:
        p=doc.add_paragraph()
        format_paragraph(p)
        p.paragraph_format.space_before=Pt(8)
        render_inline(caption,p,{"size":10,"colour":"333333"},available,source_path.parent)
        notes_and_caption.append((caption,p))
        if index<3:
            # Main-table captions precede their tables, matching the manuscript
            # display order. Keep the caption left-aligned and with the header.
            word_table._tbl.addprevious(p._p)
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(8)
    if notes_and_caption:
        # Retain one data row with its notes when a table continues, instead of
        # leaving an isolated notes-only page. No text or row order changes.
        for cell in word_table.rows[-1].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next=True
        notes_and_caption[-1][1].paragraph_format.keep_with_next=False
        if index<3 and caption is not None:
            notes_and_caption[-1][1].paragraph_format.keep_with_next=True
            if len(notes_and_caption)>1:
                notes_and_caption[-2][1].paragraph_format.keep_with_next=False
        for source,paragraph in notes_and_caption:
            if compact(visible_text(source))!=compact(paragraph.text):
                raise RuntimeError(f"Footnote or caption text changed in {key}")
    actual=[compact("".join(cell._tc.xpath(".//w:t/text()"))) for cell in destination_cells]
    if originals!=actual:
        failures=[i for i,(a,b) in enumerate(zip(originals,actual)) if a!=b]
        raise RuntimeError(f"Displayed text changed in {key}, cells {failures}")
    if key=="Table_S2" and s2_nowrap_spans_verified!=170:
        raise RuntimeError(f"S2 no-break inventory changed: {s2_nowrap_spans_verified}")
    output=outdir/(key+".docx")
    doc.save(output)
    check=Document(output)
    if len(check.tables)!=1:
        raise RuntimeError(f"{key} must contain exactly one native table")
    return {
        "label":key,"title":label,"path":str(output.resolve()),
        "sha256":hashlib.sha256(output.read_bytes()).hexdigest(),
        "rows":len(rows),"columns":ncols,"source_cells":len(originals),
        "all_displayed_cell_text_exact":True,
        "all_footnote_and_caption_text_exact":True,
        "embedded_distribution_images":len(check.inline_shapes),
        "base_font":"Arial 10 pt", "orientation":"landscape" if wide else "portrait",
        "paper":"A3" if key=="Table_S2" else "A4",
        "source_table_sha256":hashlib.sha256(html.tostring(source_table)).hexdigest(),
        **({"complete_mean_sd_nowrap_units":s2_nowrap_spans_verified,
            "numeric_base_font_pt":10-s2_numeric_reduction_pt,
            "numeric_reduction_pt":s2_numeric_reduction_pt} if key=="Table_S2" else {}),
    }


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_html",type=Path)
    parser.add_argument("reference_docx",type=Path)
    parser.add_argument("output_dir",type=Path)
    parser.add_argument("--only",nargs="+")
    parser.add_argument("--s2-numeric-reduction-pt",type=int,choices=(0,1),default=0)
    args=parser.parse_args()
    args.output_dir.mkdir(parents=True,exist_ok=True)
    root=html.parse(str(args.source_html))
    tables=root.xpath("//table[contains(concat(' ',normalize-space(@class),' '),' gt_table ')]")
    if len(tables)!=len(LABELS):
        raise RuntimeError(f"Expected {len(LABELS)} tables, found {len(tables)}")
    manifest=[]
    for i,source in enumerate(tables):
        if args.only and LABELS[i][0] not in args.only:
            continue
        record=export_one(source,args.source_html,args.reference_docx,args.output_dir,i,args.s2_numeric_reduction_pt)
        manifest.append(record)
        print(json.dumps(record,ensure_ascii=False),flush=True)
    manifest_path=args.output_dir/"table_manifest.json"
    if args.only and manifest_path.exists():
        existing={record["label"]:record for record in json.loads(manifest_path.read_text())}
        existing.update({record["label"]:record for record in manifest})
        manifest=[existing[key] for key,_ in LABELS if key in existing]
    manifest_path.write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    if len(manifest)==len(LABELS):
        required_metadata=("README.md", "table_manifest.json", "export_provenance.json")
        if not all((args.output_dir/name).is_file() for name in required_metadata):
            raise RuntimeError("The full table package requires README, manifest and export provenance")
        for record in manifest:
            target=Path(record["path"]).resolve()
            if (target.parent!=args.output_dir.resolve() or target.name!=record["label"]+".docx"
                    or hashlib.sha256(target.read_bytes()).hexdigest()!=record["sha256"]):
                raise RuntimeError("Every reused or regenerated table must be exact inside the new package directory")
        with ZipFile(args.output_dir/"editable_manuscript_tables.zip","w",ZIP_DEFLATED) as archive:
            for record in manifest:
                archive.write(record["path"],Path(record["path"]).name)
            for name in required_metadata:
                archive.write(args.output_dir/name,name)


if __name__=="__main__":
    raise SystemExit("PROSPECTIVE ONLY. Execution requires a separately sealed implementation order.")
