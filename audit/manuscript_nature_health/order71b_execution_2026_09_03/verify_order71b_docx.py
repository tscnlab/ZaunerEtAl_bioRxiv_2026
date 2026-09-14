#!/usr/bin/env python3

"""Fail-closed structural and semantic transport checks for Order 71b."""

from __future__ import annotations

import hashlib
import json
import posixpath
import re
import zipfile
from collections import Counter
from pathlib import Path
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[3]
EVIDENCE = ROOT / "audit/manuscript_nature_health/order71b_execution_2026_09_03"
RAW = EVIDENCE / "docx_candidate/ZaunerEtAl2026_NatHealth_phase3_brown_raw.docx"
CANDIDATE = EVIDENCE / "docx_candidate/ZaunerEtAl2026_NatHealth_phase3_brown_candidate.docx"
CANONICAL_PREIMAGE = (
    ROOT
    / "manuscript/R0_NatHealth/_output/ZaunerEtAl2026_NatHealth_phase3_brown.docx"
)
TABLE_MANIFEST = EVIDENCE / "word_capture/word_table_png_manifest.json"
FIGURE_MANIFEST = EVIDENCE / "word_capture/word_figure_png_manifest.json"
OUTPUT = EVIDENCE / "docx_structural_checks.json"

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_text(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


def rels_source(rels_path: str) -> str | None:
    if rels_path == "_rels/.rels":
        return None
    parent = posixpath.dirname(rels_path)
    source_dir = posixpath.dirname(parent)
    source_name = posixpath.basename(rels_path)[: -len(".rels")]
    return posixpath.join(source_dir, source_name)


def resolve_target(source: str | None, target: str) -> str:
    target_path = unquote(urlsplit(target).path)
    if target_path.startswith("/"):
        return posixpath.normpath(target_path.lstrip("/"))
    base = "" if source is None else posixpath.dirname(source)
    return posixpath.normpath(posixpath.join(base, target_path))


def package_checks(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        name_set = set(names)
        duplicate_parts = sorted(name for name, count in Counter(names).items() if count > 1)
        corrupt_member = archive.testzip()
        required_parts = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}

        graph: dict[str | None, list[str]] = {}
        missing_targets: list[dict[str, str]] = []
        duplicate_relationship_ids: list[dict[str, str]] = []
        missing_relationship_sources: list[dict[str, str]] = []
        external_relationships = 0
        for rels_path in sorted(name for name in names if name.endswith(".rels")):
            source = rels_source(rels_path)
            if source is not None and source not in name_set:
                missing_relationship_sources.append(
                    {"relationships": rels_path, "source": source},
                )
            root = ET.fromstring(archive.read(rels_path))
            relationships = list(root.findall(f"{{{REL_NS}}}Relationship"))
            ids = [item.attrib.get("Id", "") for item in relationships]
            duplicate_relationship_ids.extend(
                {"relationships": rels_path, "id": relationship_id}
                for relationship_id, count in Counter(ids).items()
                if count > 1
            )
            targets = []
            for relationship in relationships:
                if relationship.attrib.get("TargetMode") == "External":
                    external_relationships += 1
                    continue
                target = resolve_target(source, relationship.attrib["Target"])
                targets.append(target)
                if target not in name_set:
                    missing_targets.append(
                        {"relationships": rels_path, "target": target},
                    )
            graph[source] = targets

        reached: set[str] = set()
        queue = list(graph.get(None, []))
        while queue:
            current = queue.pop()
            if current in reached:
                continue
            reached.add(current)
            queue.extend(graph.get(current, []))
        content_parts = {
            name
            for name in names
            if name != "[Content_Types].xml" and not name.endswith(".rels")
        }
        unreachable_parts = sorted(content_parts - reached)

        content_types = ET.fromstring(archive.read("[Content_Types].xml"))
        defaults = {
            node.attrib["Extension"].lower(): node.attrib["ContentType"]
            for node in content_types.findall(f"{{{CONTENT_NS}}}Default")
        }
        overrides = {
            node.attrib["PartName"].lstrip("/"): node.attrib["ContentType"]
            for node in content_types.findall(f"{{{CONTENT_NS}}}Override")
        }
        uncovered_content_types = []
        for name in names:
            if name == "[Content_Types].xml":
                continue
            extension = (
                "rels"
                if name.endswith(".rels")
                else posixpath.splitext(name)[1].lstrip(".").lower()
            )
            if name not in overrides and extension not in defaults:
                uncovered_content_types.append(name)

        return {
            "part_count": len(names),
            "duplicate_parts": duplicate_parts,
            "corrupt_member": corrupt_member,
            "required_parts_present": required_parts <= name_set,
            "missing_relationship_sources": missing_relationship_sources,
            "missing_relationship_targets": missing_targets,
            "duplicate_relationship_ids": duplicate_relationship_ids,
            "external_relationship_count": external_relationships,
            "unreachable_parts": unreachable_parts,
            "uncovered_content_types": uncovered_content_types,
        }


def body_paragraphs(document) -> list:
    return [
        paragraph
        for paragraph in document.paragraphs
        if clean_text(paragraph.text)
    ]


def nested_caption_text(document, prefix: str) -> str:
    matches = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = clean_text(paragraph.text)
                    if text.startswith(prefix):
                        matches.append(text)
    if len(matches) != 1:
        raise AssertionError(f"Expected one nested caption {prefix!r}, found {len(matches)}")
    return matches[0]


def document_image_payloads(path: Path) -> tuple[list[dict[str, str]], set[str]]:
    with zipfile.ZipFile(path) as archive:
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        rels_xml = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        relationships = {
            node.attrib["Id"]: node.attrib["Target"]
            for node in rels_xml.findall(f"{{{REL_NS}}}Relationship")
        }
        namespace = {
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        payloads = []
        used_relationships = set()
        for inline in document_xml.findall(".//wp:inline", namespace):
            properties = inline.find("wp:docPr", namespace)
            blip = inline.find(".//a:blip", namespace)
            if properties is None or blip is None:
                raise AssertionError("Inline drawing lacks docPr or blip")
            relationship_id = blip.attrib[
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            ]
            used_relationships.add(relationship_id)
            target = resolve_target("word/document.xml", relationships[relationship_id])
            payloads.append(
                {
                    "alt": properties.attrib.get("descr", ""),
                    "title": properties.attrib.get("title", ""),
                    "relationship_id": relationship_id,
                    "target": target,
                    "sha256": sha256_bytes(archive.read(target)),
                },
            )
        image_relationships = {
            relationship_id
            for relationship_id, target in relationships.items()
            if target.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"))
        }
        return payloads, image_relationships - used_relationships


def style_inherits(style, base_style_id: str) -> bool:
    current = style
    while current is not None:
        if current.style_id == base_style_id:
            return True
        current = current.base_style
    return False


def ordered_numeric_tokens(texts: list[str]) -> list[str]:
    pattern = re.compile(
        r"(?<![A-Za-z0-9_])[-+]?(?:[0-9]+(?:\.[0-9]+)?|\.[0-9]+)"
        r"(?:[eE][-+]?[0-9]+)?%?",
    )
    return pattern.findall("\n".join(texts))


def paragraph_superscript_text(paragraph) -> list[str]:
    values = []
    for run in paragraph._p.xpath(
        ".//w:r[w:rPr/w:vertAlign[@w:val='superscript']]",
    ):
        text = "".join(node.text or "" for node in run.xpath(".//w:t"))
        if text:
            values.append(text)
    return values


def main() -> None:
    for path in (RAW, CANDIDATE, CANONICAL_PREIMAGE, TABLE_MANIFEST, FIGURE_MANIFEST):
        if not path.is_file():
            raise FileNotFoundError(path)

    assert sha256(CANONICAL_PREIMAGE) == "6cd592391f720829ed278a35e81c854ee7fd54eca6a4a4e664faad7b37310f91"
    candidate_package = package_checks(CANDIDATE)
    assert not candidate_package["duplicate_parts"]
    assert candidate_package["corrupt_member"] is None
    assert candidate_package["required_parts_present"]
    assert not candidate_package["missing_relationship_sources"]
    assert not candidate_package["missing_relationship_targets"]
    assert not candidate_package["duplicate_relationship_ids"]
    assert not candidate_package["unreachable_parts"]
    assert not candidate_package["uncovered_content_types"]

    raw = Document(RAW)
    candidate = Document(CANDIDATE)
    canonical = Document(CANONICAL_PREIMAGE)

    normal = candidate.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size is not None and normal.font.size.pt == 11
    for style_name in ("Body Text", "First Paragraph", "Compact", "Bibliography"):
        assert style_inherits(candidate.styles[style_name], normal.style_id)
    for style_name, size in {
        "Title": 26,
        "Heading 1": 18,
        "Heading 2": 16,
        "Heading 3": 14,
    }.items():
        assert candidate.styles[style_name].font.size.pt == size

    author_names = [
        "Johannes Zauner", "Altug Didikoglu", "Sam Aerts",
        "Gabriel Kwaku Agbeshie", "Kwadwo Owusu Akuffo", "Sena Gulsum Akgun",
        "Sema Nur Aydin", "David Baeza Moyano", "Daan Boesten", "John F.B. Bolte",
        "Kai Broszio", "Guadalupe Cantarero García", "Roberto Alonso González Lezcano",
        "Carolina Guidolin", "Sarina Hilden", "Nico Hogervorst", "Astrid Jansen",
        "Zeynep Kayar", "Stefan Källberg", "Suyoun Lee", "Sofía Melero Tur",
        "Maria Nilsson Tengelin", "María Concepción Pérez Gutiérrez",
        "Andrea Sancho-Salas", "Oliver Stefani", "Ingemar Svensson",
        "Helga von-Breymann", "Manuel Spitschan",
    ]
    candidate_paragraphs = body_paragraphs(candidate)
    raw_paragraphs = body_paragraphs(raw)
    author_matches = [
        paragraph
        for paragraph in candidate_paragraphs
        if all(name in paragraph.text for name in author_names)
    ]
    assert len(author_matches) == 1
    author_line = clean_text(author_matches[0].text)
    positions = [author_line.index(name) for name in author_names]
    assert positions == sorted(positions)
    assert all(author_line.count(name) == 1 for name in author_names)
    assert clean_text(raw_paragraphs[1].text) == author_line

    affiliation_text = next(
        clean_text(paragraph.text)
        for paragraph in raw_paragraphs
        if clean_text(paragraph.text).startswith(
            "1 Department Health and Sports Sciences",
        )
    )
    correspondence_text = "✉ Correspondence: Johannes Zauner <johannes.zauner@tum.de>"
    assert sum(clean_text(p.text) == affiliation_text for p in candidate_paragraphs) == 1
    assert sum(clean_text(p.text) == correspondence_text for p in candidate_paragraphs) == 1
    assert sum(re.fullmatch(r"\d{4}-\d{2}-\d{2}", clean_text(p.text)) is not None for p in candidate_paragraphs) == 1
    assert sum(clean_text(p.text) == "Abstract" for p in candidate_paragraphs) == 1

    body_children = list(candidate.element.body)
    front_elements = {}
    for key, text in {
        "date": "2026-08-21",
        "affiliations": affiliation_text,
        "correspondence": correspondence_text,
        "abstract": "Abstract",
    }.items():
        matches = [
            paragraph
            for paragraph in candidate_paragraphs
            if clean_text(paragraph.text) == text
        ]
        assert len(matches) == 1
        front_elements[key] = body_children.index(matches[0]._p)
    assert list(front_elements.values()) == [2, 3, 4, 5]

    raw_affiliation = next(p for p in raw_paragraphs if clean_text(p.text) == affiliation_text)
    raw_correspondence = next(
        p for p in raw_paragraphs if clean_text(p.text) == correspondence_text
    )
    candidate_affiliation = next(
        p for p in candidate_paragraphs if clean_text(p.text) == affiliation_text
    )
    candidate_correspondence = next(
        p for p in candidate_paragraphs if clean_text(p.text) == correspondence_text
    )
    assert candidate_affiliation._p.xml == raw_affiliation._p.xml
    assert candidate_correspondence._p.xml == raw_correspondence._p.xml
    assert candidate_affiliation.style.style_id == raw_affiliation.style.style_id
    assert candidate_correspondence.style.style_id == raw_correspondence.style.style_id

    raw_ordered_paragraphs = list(raw_paragraphs)
    raw_ordered_paragraphs.remove(raw_affiliation)
    raw_ordered_paragraphs.remove(raw_correspondence)
    date_position = next(
        index
        for index, paragraph in enumerate(raw_ordered_paragraphs)
        if clean_text(paragraph.text) == "2026-08-21"
    )
    raw_ordered_paragraphs[date_position + 1 : date_position + 1] = [
        raw_affiliation,
        raw_correspondence,
    ]
    raw_sequence = [clean_text(paragraph.text) for paragraph in raw_ordered_paragraphs]
    remaining = Counter(raw_sequence)
    preserved_candidate_sequence = []
    preserved_candidate_paragraphs = []
    extras = []
    for paragraph in candidate_paragraphs:
        text = clean_text(paragraph.text)
        if remaining[text] > 0:
            preserved_candidate_sequence.append(text)
            preserved_candidate_paragraphs.append(paragraph)
            remaining[text] -= 1
        else:
            extras.append(text)
    assert sum(remaining.values()) == 0
    assert preserved_candidate_sequence == raw_sequence
    assert ordered_numeric_tokens(preserved_candidate_sequence) == ordered_numeric_tokens(raw_sequence)
    assert [
        paragraph_superscript_text(paragraph)
        for paragraph in preserved_candidate_paragraphs
    ] == [
        paragraph_superscript_text(paragraph)
        for paragraph in raw_ordered_paragraphs
    ]

    nested_captions = {
        f"Figure {number}": nested_caption_text(raw, f"Figure {number}:")
        for number in range(1, 4)
    }
    nested_captions.update(
        {
            f"Table {number}": nested_caption_text(raw, f"Table {number}:")
            for number in range(1, 4)
        },
    )
    added_labels = {
        "Supplementary Figure S8 (continued)",
        "Supplementary Table S11. Light-exposure behaviour and awareness",
        "Supplementary Table S11 (continued).",
        "Supplementary Table S12. Visual light sensitivity",
        "Supplementary Table S13. Chronotype and timing",
        "Supplementary Table S14. Age and biological sex",
        "Supplementary Table S15. Biological-sex-specific daily patterns and global tests",
    }
    assert Counter(extras) == Counter([*nested_captions.values(), *added_labels])
    for caption in nested_captions.values():
        assert sum(clean_text(p.text) == caption for p in candidate_paragraphs) == 1

    assert len(candidate.sections) == 27
    expected_orientations = [
        orientation
        for pair in ((WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE),) * 13
        for orientation in pair
    ] + [WD_ORIENT.PORTRAIT]
    assert [section.orientation for section in candidate.sections] == expected_orientations
    for index, section in enumerate(candidate.sections):
        landscape = index % 2 == 1
        if landscape:
            assert (section.page_width, section.page_height) == (10693400, 7557135)
            assert (section.left_margin, section.right_margin) == (411480, 411480)
            assert (section.top_margin, section.bottom_margin) == (502920, 502920)
        else:
            assert (section.page_width, section.page_height) == (7557135, 10693400)
            assert (section.left_margin, section.right_margin) == (914400, 914400)
            assert (section.top_margin, section.bottom_margin) == (914400, 914400)
        line_numbers = section._sectPr.findall(qn("w:lnNumType"))
        assert len(line_numbers) == 1
        assert line_numbers[0].get(qn("w:countBy")) == "1"
        assert line_numbers[0].get(qn("w:restart")) == "continuous"
        if landscape:
            assert line_numbers[0].get(qn("w:distance")) == "72"

    section_breaks = candidate.element.body.xpath("./w:p[w:pPr/w:sectPr]")
    assert len(section_breaks) == 26
    assert all(
        paragraph.xpath("./w:pPr/w:suppressLineNumbers")
        for paragraph in section_breaks
    )

    footer_parts = {}
    for footer_reference in candidate.element.body.xpath(
        ".//w:sectPr/w:footerReference",
    ):
        relationship_id = footer_reference.get(qn("r:id"))
        part = candidate.part.related_parts[relationship_id]
        footer_parts[str(part.partname)] = part
    assert len(footer_parts) == 2
    footer_page_fields = 0
    for footer_part in footer_parts.values():
        footer_paragraphs = footer_part.element.xpath(".//w:p")
        assert len(footer_paragraphs) == 2
        assert all(
            paragraph.xpath("./w:pPr/w:suppressLineNumbers")
            for paragraph in footer_paragraphs
        )
        instructions = [
            clean_text(node.text or "").upper()
            for node in footer_part.element.xpath(".//w:instrText")
        ]
        simple = [
            clean_text(node.get(qn("w:instr"), "")).upper()
            for node in footer_part.element.xpath(".//w:fldSimple")
        ]
        assert instructions.count("PAGE") + simple.count("PAGE") == 1
        footer_page_fields += 1
    assert footer_page_fields == 2
    assert len(candidate.settings._element.xpath("./w:doNotAutoCompressPictures")) == 1

    assert len(candidate.tables) == 0
    assert len(candidate.inline_shapes) == 53
    properties = candidate.element.body.xpath(".//wp:docPr")
    assert len(properties) == 53
    assert all(clean_text(node.get("descr", "")) for node in properties)
    assert all(clean_text(node.get("title", "")) for node in properties)
    extents = candidate.element.body.xpath(".//wp:inline/wp:extent")
    assert len(extents) == 53
    assert all(int(node.get("cx")) > 0 and int(node.get("cy")) > 0 for node in extents)
    crops = [dict(node.attrib) for node in candidate.element.body.xpath(".//a:srcRect")]
    assert crops == [{"b": "57695"}, {"t": "42305"}]
    alt_values = [node.get("descr") for node in properties]
    assert alt_values.count("Supplementary Figure S8, panels A-C") == 1
    assert alt_values.count("Supplementary Figure S8, panel D") == 1

    table_manifest = json.loads(TABLE_MANIFEST.read_text(encoding="utf-8"))
    figure_manifest = json.loads(FIGURE_MANIFEST.read_text(encoding="utf-8"))
    expected_table_hashes = Counter(
        sha256(Path(item["path"]))
        for entry in table_manifest
        for item in entry["files"]
    )
    expected_figure_hashes = Counter(
        sha256(Path(entry["path"]))
        for entry in figure_manifest
        for _ in range(2 if entry["key"] == "supp_figure_s8" else 1)
    )
    candidate_payloads, unused_image_relationships = document_image_payloads(CANDIDATE)
    canonical_payloads, canonical_unused_image_relationships = document_image_payloads(
        CANONICAL_PREIMAGE,
    )
    embedded_table_hashes = Counter(
        item["sha256"]
        for item in candidate_payloads
        if item["alt"].startswith(("Main Table ", "Supplementary Table "))
    )
    embedded_supplementary_figure_hashes = Counter(
        item["sha256"]
        for item in candidate_payloads
        if item["alt"].startswith("Supplementary Figure ")
    )
    assert embedded_table_hashes == expected_table_hashes
    assert embedded_supplementary_figure_hashes == expected_figure_hashes
    main_figure_hashes = {
        item["alt"]: item["sha256"]
        for item in candidate_payloads
        if item["alt"].startswith("Main Figure ")
    }
    canonical_main_figure_hashes = {
        item["alt"]: item["sha256"]
        for item in canonical_payloads
        if item["alt"].startswith("Main Figure ")
    }
    assert main_figure_hashes == canonical_main_figure_hashes
    assert len(unused_image_relationships) == 17
    assert len(canonical_unused_image_relationships) == 17

    bookmarks = {
        node.get(qn("w:name"))
        for node in candidate.element.body.xpath(".//w:bookmarkStart")
    }
    internal_anchors = [
        node.get(qn("w:anchor"))
        for node in candidate.element.body.xpath(".//w:hyperlink[@w:anchor]")
    ]
    assert all(anchor in bookmarks for anchor in internal_anchors)

    candidate_texts = [clean_text(paragraph.text) for paragraph in candidate.paragraphs]
    for heading in (
        "The multiscale architecture of personal light exposure",
        "Abstract", "Results", "Discussion", "Methods", "Data availability",
        "Code availability", "References", "Acknowledgements", "Funding",
        "Author contributions", "Competing interests", "Supplementary Information",
    ):
        assert candidate_texts.count(heading) == 1
    assert sum(paragraph.style.name == "Bibliography" for paragraph in candidate.paragraphs) == 91
    for number in range(1, 4):
        assert sum(text.startswith(f"Figure {number}:") for text in candidate_texts) == 1
        assert sum(text.startswith(f"Table {number}:") for text in candidate_texts) == 1
    supplementary_identity_indices = set()

    def require_supplementary_identity(prefix, expected_styles):
        matches = [
            (index, paragraph)
            for index, paragraph in enumerate(candidate.paragraphs)
            if clean_text(paragraph.text).startswith(prefix)
        ]
        assert Counter(paragraph.style.name for _, paragraph in matches) == Counter(
            expected_styles,
        )
        supplementary_identity_indices.update(index for index, _ in matches)

    for number in range(1, 13):
        require_supplementary_identity(
            f"Supplementary Figure S{number}.",
            ["Heading 3", "Body Text"],
        )
    require_supplementary_identity(
        "Supplementary Figure S8 (continued)",
        ["Caption"],
    )
    for number in range(13, 18):
        require_supplementary_identity(
            f"Supplementary Figure S{number} and Table S{number - 2}.",
            ["Heading 3"],
        )
        require_supplementary_identity(
            f"Supplementary Figure S{number}.",
            ["Body Text"],
        )
    for number in range(1, 11):
        require_supplementary_identity(
            f"Supplementary Table S{number}.",
            ["Heading 3", "Body Text"],
        )
    for number in range(11, 16):
        require_supplementary_identity(
            f"Supplementary Table S{number}.",
            ["Caption"],
        )
    require_supplementary_identity(
        "Supplementary Table S11 (continued).",
        ["Caption"],
    )
    all_supplementary_identity_indices = {
        index
        for index, text in enumerate(candidate_texts)
        if text.startswith(("Supplementary Figure S", "Supplementary Table S"))
    }
    assert len(supplementary_identity_indices) == 61
    assert supplementary_identity_indices == all_supplementary_identity_indices

    results = {
        "status": "PASS",
        "candidate": {
            "path": str(CANDIDATE),
            "sha256": sha256(CANDIDATE),
            "bytes": CANDIDATE.stat().st_size,
        },
        "raw_docx": {
            "path": str(RAW),
            "sha256": sha256(RAW),
            "bytes": RAW.stat().st_size,
        },
        "canonical_preimage": {
            "path": str(CANONICAL_PREIMAGE),
            "sha256": sha256(CANONICAL_PREIMAGE),
            "bytes": CANONICAL_PREIMAGE.stat().st_size,
        },
        "package": candidate_package,
        "author_block": {
            "compact_author_paragraphs": len(author_matches),
            "authors_in_order": len(author_names),
            "affiliations_in_one_paragraph": 14,
            "front_body_indices": front_elements,
            "affiliation_xml_preserved": True,
            "correspondence_xml_preserved": True,
        },
        "protected_transport": {
            "raw_top_level_nonempty_paragraphs": len(raw_sequence),
            "candidate_preserved_paragraphs": len(preserved_candidate_sequence),
            "candidate_expected_added_labels_and_captions": len(extras),
            "ordered_numeric_tokens": len(ordered_numeric_tokens(raw_sequence)),
            "superscript_runs": sum(
                len(paragraph_superscript_text(paragraph))
                for paragraph in raw_ordered_paragraphs
            ),
            "all_text_citations_and_ordered_numbers_preserved": True,
        },
        "styles": {
            "normal_font": normal.font.name,
            "normal_size_pt": normal.font.size.pt,
            "body_styles_inherit_normal": True,
        },
        "sections": {
            "count": len(candidate.sections),
            "portrait": sum(s.orientation == WD_ORIENT.PORTRAIT for s in candidate.sections),
            "landscape": sum(s.orientation == WD_ORIENT.LANDSCAPE for s in candidate.sections),
            "section_breaks_suppress_line_numbers": len(section_breaks),
            "continuous_line_numbering": True,
            "footer_parts": len(footer_parts),
            "page_fields": footer_page_fields,
        },
        "displays": {
            "native_tables": len(candidate.tables),
            "inline_shapes": len(candidate.inline_shapes),
            "alt_text_entries": len(properties),
            "crop_rectangles": crops,
            "fresh_table_pngs_embedded_exact": sum(expected_table_hashes.values()),
            "fresh_supplementary_figure_png_drawings_embedded_exact": sum(
                expected_figure_hashes.values()
            ),
            "main_figures_match_protected_preimage": len(main_figure_hashes),
        },
        "references": {
            "bibliography_entries": 91,
            "internal_hyperlink_anchors": len(internal_anchors),
            "all_internal_anchors_resolve": True,
        },
        "informational": {
            "unused_document_image_relationships": len(unused_image_relationships),
            "protected_preimage_unused_document_image_relationships": len(
                canonical_unused_image_relationships
            ),
            "note": (
                "These retained image relationships are package-reachable and have valid "
                "targets. Their count is unchanged from the independently accepted preimage."
            ),
        },
    }
    OUTPUT.write_text(json.dumps(results, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(results, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
