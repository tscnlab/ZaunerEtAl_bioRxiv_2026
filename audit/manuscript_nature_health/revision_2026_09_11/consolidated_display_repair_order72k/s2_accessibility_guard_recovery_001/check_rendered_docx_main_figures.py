"""Read-only structural precondition check for the frozen Word assembler."""
import json
from pathlib import Path
import sys
from docx import Document

source = Path(sys.argv[1])
doc = Document(source)
records = []
for number in range(1, 4):
    wrappers = [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1
                and " ".join(t.cell(0, 0).text.split()).startswith(f"Figure {number}:")]
    entry = {"figure": number, "wrapper_count": len(wrappers)}
    if len(wrappers) == 1:
        ps = wrappers[0].cell(0, 0).paragraphs
        images = [p for p in ps if p._p.xpath(".//w:drawing")]
        captions = [p for p in ps if " ".join(p.text.split()).startswith(f"Figure {number}:")]
        entry.update(image_paragraph_count=len(images), caption_paragraph_count=len(captions),
                     paragraph_texts=[p.text for p in ps], drawings=[])
        for p in images:
            for blip in p._p.xpath(".//a:blip"):
                rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                rel = doc.part.rels.get(rid)
                entry["drawings"].append({"relationship": rid,
                    "target": str(rel.target_part.partname) if rel is not None and not rel.is_external else None,
                    "content_type": rel.target_part.content_type if rel is not None and not rel.is_external else None})
        entry["assembler_precondition_pass"] = len(images) == 1 and len(captions) == 1
    else:
        entry["assembler_precondition_pass"] = False
    records.append(entry)
print(json.dumps({"source": str(source), "main_figures": records,
                  "all_preconditions_pass": all(r["assembler_precondition_pass"] for r in records)}, indent=2))
