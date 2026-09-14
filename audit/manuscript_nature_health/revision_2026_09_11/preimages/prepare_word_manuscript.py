#!/usr/bin/env python3

"""Replace lossy gt-to-DOCX tables with faithful PNGs and tune page layout."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path

from lxml import etree
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("png_manifest", type=Path)
    parser.add_argument("figure_manifest", type=Path)
    parser.add_argument("output_docx", type=Path)
    return parser.parse_args()


def local_name(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def top_level_paragraphs(document: Document) -> list[Paragraph]:
    return [
        Paragraph(element, document)
        for element in document.element.body
        if local_name(element) == "p"
    ]


def top_level_tables(document: Document) -> list[Table]:
    return [
        Table(element, document)
        for element in document.element.body
        if local_name(element) == "tbl"
    ]


def clean_text(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


EXPECTED_AFFILIATIONS = (
    "Department Health and Sports Sciences, TUM School of Medicine and Health, "
    "Chronobiology & Health, Technical University of Munich, Munich, Germany",
    "Division of Diabetes, Endocrinology, and Gastroenterology, School of Medical "
    "Sciences, Faculty of Biology, Medicine and Health, The University of "
    "Manchester, Manchester, UK",
    "Faculty of Science, Department of Neuroscience, Izmir Institute of "
    "Technology, Izmir, Türkiye",
    "Smart Sensor Systems research group, The Hague University of Applied "
    "Sciences, Delft, The Netherlands",
    "Department of Optometry and Visual Science, College of Science, Kwame "
    "Nkrumah University of Science and Technology, Kumasi, Ghana",
    "Fundación Universitaria CEU San Pablo, Boadilla del Monte, Madrid, Spain",
    "Centre for Sustainability, Environment and Health, National Institute for "
    "Public Health and the Environment (RIVM), Bilthoven, The Netherlands",
    "Federal Institute for Occupational Safety and Health (BAuA), Dortmund, Germany",
    "Max Planck Institute for Biological Cybernetics, Max Planck Research Group "
    "Translational Sensory & Circadian Neuroscience, Tübingen, Germany",
    "RISE Research Institutes of Sweden, Borås, Sweden",
    "School of Architecture, University of Costa Rica, San José, Costa Rica",
    "Lucerne University of Applied Sciences and Arts, Lucerne, Switzerland",
    "TUM Institute for Advanced Study (TUM-IAS), Technical University of Munich, "
    "Garching, Germany",
    "TUMCREATE Ltd., Singapore, Singapore",
)
EXPECTED_AFFILIATION_TEXT = clean_text(
    " ".join(
        f"{number} {affiliation}"
        for number, affiliation in enumerate(EXPECTED_AFFILIATIONS, start=1)
    ),
)
EXPECTED_CORRESPONDENCE_TEXT = (
    "✉ Correspondence: Johannes Zauner <johannes.zauner@tum.de>"
)
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

MAIN_DISPLAY_BOOKMARK_DESTINATIONS = (
    ("fig-study-overview", "Figure 1:"),
    ("tbl-participant-site", "Table 1:"),
    ("tbl-brown-adherence", "Table 2:"),
    ("fig-daily-architecture", "Figure 2:"),
    ("tbl-metric-context", "Table 3:"),
    ("fig-activity-context", "Figure 3:"),
)
SUPPLEMENTARY_DISPLAY_BOOKMARK_NUMBERS = tuple(
    number for number in range(1, 18) if number != 3
)
EXPECTED_MISSING_DISPLAY_BOOKMARKS = frozenset(
    {
        *(name for name, _ in MAIN_DISPLAY_BOOKMARK_DESTINATIONS),
        *(f"fig-s{number}" for number in SUPPLEMENTARY_DISPLAY_BOOKMARK_NUMBERS),
    },
)


def relationship_signature(document: Document) -> tuple[tuple[str, str, str, bool], ...]:
    return tuple(
        sorted(
            (
                relationship_id,
                relationship.reltype,
                relationship.target_ref,
                relationship.is_external,
            )
            for relationship_id, relationship in document.part.rels.items()
        ),
    )


def add_missing_display_bookmarks(document: Document) -> dict[str, object]:
    root = document.element
    relationship_signature_before = relationship_signature(document)
    document_xml_before = root.xml
    text_before = tuple(element.text or "" for element in root.xpath(".//w:t"))

    hyperlinks_before = root.xpath(".//w:hyperlink[@w:anchor]")
    anchors_before = tuple(
        hyperlink.get(qn("w:anchor")) for hyperlink in hyperlinks_before
    )
    if len(anchors_before) != 124 or len(set(anchors_before)) != 102:
        raise RuntimeError(
            "Display-bookmark repair requires 124 internal hyperlinks and "
            f"102 unique targets; found {len(anchors_before)} and "
            f"{len(set(anchors_before))}",
        )
    if "fig-s3" in anchors_before:
        raise RuntimeError("fig-s3 must remain absent from internal hyperlink targets")

    starts_before = root.xpath(".//w:bookmarkStart")
    ends_before = root.xpath(".//w:bookmarkEnd")
    start_names_before = tuple(
        bookmark.get(qn("w:name")) for bookmark in starts_before
    )
    start_ids_before = tuple(bookmark.get(qn("w:id")) for bookmark in starts_before)
    end_ids_before = tuple(bookmark.get(qn("w:id")) for bookmark in ends_before)
    if (
        len(starts_before) != 173
        or len(set(start_names_before)) != 173
        or len(set(start_ids_before)) != 173
        or len(ends_before) != 173
        or len(set(end_ids_before)) != 173
        or set(start_ids_before) != set(end_ids_before)
    ):
        raise RuntimeError(
            "Display-bookmark repair requires 173 unique paired bookmark names "
            "and IDs before mutation",
        )
    if "fig-s3" in start_names_before:
        raise RuntimeError("fig-s3 must remain absent from bookmark names")
    try:
        numeric_ids_before = tuple(int(value) for value in start_ids_before)
    except (TypeError, ValueError) as exc:
        raise RuntimeError("Every existing bookmark ID must be numeric") from exc

    missing_before = frozenset(set(anchors_before) - set(start_names_before))
    if missing_before != EXPECTED_MISSING_DISPLAY_BOOKMARKS:
        raise RuntimeError(
            "Unexpected missing internal targets before bookmark repair: "
            f"{sorted(missing_before)}",
        )

    protected_starts = {
        (bookmark.get(qn("w:id")), bookmark.get(qn("w:name"))): etree.tostring(
            bookmark,
            encoding="unicode",
        )
        for bookmark in starts_before
    }
    protected_ends = {
        bookmark.get(qn("w:id")): etree.tostring(bookmark, encoding="unicode")
        for bookmark in ends_before
    }

    destinations: dict[str, Paragraph] = {}
    for target, prefix in MAIN_DISPLAY_BOOKMARK_DESTINATIONS:
        matches = find_paragraphs(document, prefix)
        if len(matches) != 1:
            raise RuntimeError(
                f"Expected one top-level caption beginning {prefix!r}, "
                f"found {len(matches)}",
            )
        destinations[target] = matches[0]

    for number in SUPPLEMENTARY_DISPLAY_BOOKMARK_NUMBERS:
        pattern = re.compile(
            rf"^Supplementary Figure S{number}(?:\.| and Table\b)",
        )
        matches = [
            paragraph
            for paragraph in top_level_paragraphs(document)
            if paragraph.style.name == "Heading 3"
            and pattern.match(clean_text(paragraph.text))
        ]
        if len(matches) != 1:
            raise RuntimeError(
                "Expected one Heading 3 destination for Supplementary Figure "
                f"S{number}; found {len(matches)}",
            )
        destinations[f"fig-s{number}"] = matches[0]

    if set(destinations) != EXPECTED_MISSING_DISPLAY_BOOKMARKS:
        raise RuntimeError(
            "Resolved display destinations differ from the sealed 22-name set",
        )
    if len({id(paragraph._p) for paragraph in destinations.values()}) != 22:
        raise RuntimeError("Every display bookmark must have a unique destination")

    next_id = max(numeric_ids_before) + 1
    added = []
    for offset, (target, paragraph) in enumerate(destinations.items()):
        bookmark_id = str(next_id + offset)
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), bookmark_id)
        bookmark_start.set(qn("w:name"), target)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), bookmark_id)

        paragraph_properties = paragraph._p.find(qn("w:pPr"))
        insertion_index = 1 if paragraph_properties is not None else 0
        paragraph._p.insert(insertion_index, bookmark_start)
        paragraph._p.insert(insertion_index + 1, bookmark_end)
        added.append(
            {
                "name": target,
                "id": bookmark_id,
                "destination": clean_text(paragraph.text),
            },
        )

    hyperlinks_after = root.xpath(".//w:hyperlink[@w:anchor]")
    anchors_after = tuple(
        hyperlink.get(qn("w:anchor")) for hyperlink in hyperlinks_after
    )
    starts_after = root.xpath(".//w:bookmarkStart")
    ends_after = root.xpath(".//w:bookmarkEnd")
    start_names_after = tuple(
        bookmark.get(qn("w:name")) for bookmark in starts_after
    )
    start_ids_after = tuple(bookmark.get(qn("w:id")) for bookmark in starts_after)
    end_ids_after = tuple(bookmark.get(qn("w:id")) for bookmark in ends_after)

    if anchors_after != anchors_before:
        raise RuntimeError("Internal hyperlinks changed during bookmark repair")
    if "fig-s3" in anchors_after or "fig-s3" in start_names_after:
        raise RuntimeError("fig-s3 must remain absent after bookmark repair")
    if (
        len(starts_after) != 195
        or len(set(start_names_after)) != 195
        or len(set(start_ids_after)) != 195
        or len(ends_after) != 195
        or len(set(end_ids_after)) != 195
        or set(start_ids_after) != set(end_ids_after)
    ):
        raise RuntimeError(
            "Display-bookmark repair did not produce 195 unique paired "
            "bookmark names and IDs",
        )
    if set(anchors_after) - set(start_names_after):
        raise RuntimeError("One or more internal hyperlinks remain unresolved")

    starts_after_map = {
        (bookmark.get(qn("w:id")), bookmark.get(qn("w:name"))): etree.tostring(
            bookmark,
            encoding="unicode",
        )
        for bookmark in starts_after
    }
    ends_after_map = {
        bookmark.get(qn("w:id")): etree.tostring(bookmark, encoding="unicode")
        for bookmark in ends_after
    }
    if any(
        starts_after_map.get(key) != value for key, value in protected_starts.items()
    ):
        raise RuntimeError("A pre-existing bookmark start changed during repair")
    if any(ends_after_map.get(key) != value for key, value in protected_ends.items()):
        raise RuntimeError("A pre-existing bookmark end changed during repair")
    if relationship_signature(document) != relationship_signature_before:
        raise RuntimeError("Document relationships changed during bookmark repair")
    if tuple(element.text or "" for element in root.xpath(".//w:t")) != text_before:
        raise RuntimeError("Visible document text changed during bookmark repair")

    reversed_root = deepcopy(root)
    added_ids = {entry["id"] for entry in added}
    for bookmark in list(reversed_root.xpath(".//w:bookmarkStart")):
        if bookmark.get(qn("w:name")) in EXPECTED_MISSING_DISPLAY_BOOKMARKS:
            bookmark.getparent().remove(bookmark)
    for bookmark in list(reversed_root.xpath(".//w:bookmarkEnd")):
        if bookmark.get(qn("w:id")) in added_ids:
            bookmark.getparent().remove(bookmark)
    if reversed_root.xml != document_xml_before:
        raise RuntimeError(
            "Removing the 22 new bookmark pairs did not exactly reproduce the "
            "pre-repair document XML",
        )

    return {
        "internal_hyperlinks": len(anchors_after),
        "unique_internal_targets": len(set(anchors_after)),
        "bookmark_names_before": len(set(start_names_before)),
        "bookmark_names_after": len(set(start_names_after)),
        "unresolved_before": sorted(missing_before),
        "unresolved_after": [],
        "fig_s3_hyperlinks": anchors_after.count("fig-s3"),
        "fig_s3_bookmarks": start_names_after.count("fig-s3"),
        "added": added,
        "exact_reversal": True,
    }


def relocate_author_block(document: Document) -> dict[str, object]:
    paragraphs = top_level_paragraphs(document)
    affiliation_matches = [
        paragraph
        for paragraph in paragraphs
        if clean_text(paragraph.text) == EXPECTED_AFFILIATION_TEXT
    ]
    correspondence_matches = [
        paragraph
        for paragraph in paragraphs
        if clean_text(paragraph.text) == EXPECTED_CORRESPONDENCE_TEXT
    ]
    date_matches = [
        paragraph
        for paragraph in paragraphs
        if ISO_DATE.fullmatch(clean_text(paragraph.text))
    ]
    abstract_matches = [
        paragraph
        for paragraph in paragraphs
        if clean_text(paragraph.text) == "Abstract"
    ]
    match_counts = {
        "affiliations": len(affiliation_matches),
        "correspondence": len(correspondence_matches),
        "date": len(date_matches),
        "abstract": len(abstract_matches),
    }
    if set(match_counts.values()) != {1}:
        raise RuntimeError(
            "Author-block relocation requires one exact match for every anchor; "
            f"found {match_counts}",
        )

    affiliation = affiliation_matches[0]
    correspondence = correspondence_matches[0]
    date = date_matches[0]
    abstract = abstract_matches[0]
    body = document.element.body
    before_children = list(body)

    def body_index(paragraph: Paragraph) -> int:
        return before_children.index(paragraph._p)

    raw_indices = {
        "date": body_index(date),
        "abstract": body_index(abstract),
        "affiliations": body_index(affiliation),
        "correspondence": body_index(correspondence),
    }
    if not (
        raw_indices["date"]
        < raw_indices["abstract"]
        < raw_indices["affiliations"]
        < raw_indices["correspondence"]
    ):
        raise RuntimeError(f"Unexpected raw author-block order: {raw_indices}")
    if raw_indices["correspondence"] != raw_indices["affiliations"] + 1:
        raise RuntimeError("Raw affiliations and correspondence are not consecutive")

    protected = {
        "affiliations": {
            "text": affiliation.text,
            "style_id": affiliation.style.style_id,
            "xml": affiliation._p.xml,
        },
        "correspondence": {
            "text": correspondence.text,
            "style_id": correspondence.style.style_id,
            "xml": correspondence._p.xml,
        },
    }
    relationships_before = relationship_signature(document)

    affiliation_element = affiliation._p
    correspondence_element = correspondence._p
    remaining = [
        element
        for element in before_children
        if element is not affiliation_element and element is not correspondence_element
    ]
    date_position = remaining.index(date._p)
    expected_children = (
        remaining[: date_position + 1]
        + [affiliation_element, correspondence_element]
        + remaining[date_position + 1 :]
    )

    body.remove(affiliation_element)
    body.remove(correspondence_element)
    insertion_position = list(body).index(date._p) + 1
    body.insert(insertion_position, affiliation_element)
    body.insert(insertion_position + 1, correspondence_element)

    after_children = list(body)
    if len(after_children) != len(before_children) or any(
        actual is not expected
        for actual, expected in zip(after_children, expected_children, strict=True)
    ):
        raise RuntimeError("Author-block relocation changed another body element")

    final_indices = {
        "date": after_children.index(date._p),
        "affiliations": after_children.index(affiliation_element),
        "correspondence": after_children.index(correspondence_element),
        "abstract": after_children.index(abstract._p),
    }
    if list(final_indices.values()) != list(
        range(final_indices["date"], final_indices["date"] + 4),
    ):
        raise RuntimeError(f"Author-block front order is not consecutive: {final_indices}")

    for label, paragraph in (
        ("affiliations", affiliation),
        ("correspondence", correspondence),
    ):
        if paragraph.text != protected[label]["text"]:
            raise RuntimeError(f"{label.title()} text changed during relocation")
        if paragraph.style.style_id != protected[label]["style_id"]:
            raise RuntimeError(f"{label.title()} style changed during relocation")
        if paragraph._p.xml != protected[label]["xml"]:
            raise RuntimeError(f"{label.title()} XML changed during relocation")
    if relationship_signature(document) != relationships_before:
        raise RuntimeError("Document relationships changed during author-block relocation")

    return {
        "affiliations": len(EXPECTED_AFFILIATIONS),
        "raw_indices": raw_indices,
        "final_indices": final_indices,
        "affiliation_xml_sha256": hashlib.sha256(
            protected["affiliations"]["xml"].encode("utf-8"),
        ).hexdigest(),
        "correspondence_xml_sha256": hashlib.sha256(
            protected["correspondence"]["xml"].encode("utf-8"),
        ).hexdigest(),
        "relationship_count": len(relationships_before),
    }


def find_paragraphs(document: Document, prefix: str) -> list[Paragraph]:
    return [
        paragraph
        for paragraph in top_level_paragraphs(document)
        if clean_text(paragraph.text).startswith(prefix)
    ]


def find_exact_heading(document: Document, text: str) -> Paragraph:
    matches = [
        paragraph
        for paragraph in top_level_paragraphs(document)
        if clean_text(paragraph.text) == text
        and paragraph.style.name.lower().startswith("heading")
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one heading {text!r}, found {len(matches)}")
    return matches[0]


def add_picture_paragraph(
    document: Document,
    image_path: Path,
    max_width_inches: float,
    max_height_inches: float,
    alt_text: str,
    page_break_before: bool = False,
):
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    ratio = pixel_width / pixel_height
    width = min(max_width_inches, max_height_inches * ratio)
    height = width / ratio

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width), height=Inches(height))
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", alt_text)

    body = document.element.body
    body.remove(paragraph._p)
    return paragraph._p


def add_cropped_picture_paragraph(
    document: Document,
    image_path: Path,
    max_width_inches: float,
    max_height_inches: float,
    alt_text: str,
    crop_top: float = 0.0,
    crop_bottom: float = 0.0,
    keep_with_next: bool = False,
):
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    visible_fraction = 1.0 - crop_top - crop_bottom
    if not 0.0 < visible_fraction <= 1.0:
        raise ValueError("The cropped image must retain a positive height")
    ratio = pixel_width / (pixel_height * visible_fraction)
    width = min(max_width_inches, max_height_inches * ratio)
    height = width / ratio

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width), height=Inches(height))

    blip_fills = paragraph._p.xpath(".//pic:blipFill")
    if len(blip_fills) != 1:
        raise RuntimeError(
            f"Expected one picture fill for {alt_text}, found {len(blip_fills)}",
        )
    source_rectangle = OxmlElement("a:srcRect")
    if crop_top:
        source_rectangle.set("t", str(round(crop_top * 100000)))
    if crop_bottom:
        source_rectangle.set("b", str(round(crop_bottom * 100000)))
    blip_fills[0].insert(1, source_rectangle)

    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", alt_text)

    body = document.element.body
    body.remove(paragraph._p)
    return paragraph._p


def resize_drawing_paragraph(
    paragraph: Paragraph,
    max_width_inches: float,
    max_height_inches: float,
    alt_text: str,
):
    extents = paragraph._p.xpath(".//wp:inline/wp:extent | .//wp:anchor/wp:extent")
    if len(extents) != 1:
        raise RuntimeError(
            f"Expected one drawing extent for {alt_text}, found {len(extents)}",
        )
    extent = extents[0]
    width = int(extent.get("cx"))
    height = int(extent.get("cy"))
    max_width = round(max_width_inches * 914400)
    max_height = round(max_height_inches * 914400)
    scale = min(max_width / width, max_height / height, 1.0)
    new_width = round(width * scale)
    new_height = round(height * scale)
    extent.set("cx", str(new_width))
    extent.set("cy", str(new_height))
    for transform_extent in paragraph._p.xpath(".//a:xfrm/a:ext"):
        transform_extent.set("cx", str(new_width))
        transform_extent.set("cy", str(new_height))
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", alt_text)


def add_label_paragraph(
    document: Document,
    text: str,
    page_break_before: bool,
):
    paragraph = document.add_paragraph()
    paragraph.style = "Caption" if "Caption" in document.styles else paragraph.style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10)

    body = document.element.body
    body.remove(paragraph._p)
    return paragraph._p


def base_section_properties(document: Document):
    section_properties = document.element.body.find(qn("w:sectPr"))
    if section_properties is None:
        raise RuntimeError("DOCX body has no final section properties")
    return deepcopy(section_properties)


def section_break_paragraph(base_sect_pr, landscape: bool):
    sect_pr = deepcopy(base_sect_pr)
    section_type = sect_pr.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        sect_pr.append(section_type)
    section_type.set(qn("w:val"), "nextPage")

    page_size = sect_pr.find(qn("w:pgSz"))
    if page_size is None:
        page_size = OxmlElement("w:pgSz")
        sect_pr.append(page_size)
        page_width = round(8.5 * 1440)
        page_height = round(11 * 1440)
    else:
        page_width = int(page_size.get(qn("w:w")))
        page_height = int(page_size.get(qn("w:h")))
    margins = sect_pr.find(qn("w:pgMar"))
    if margins is None:
        margins = OxmlElement("w:pgMar")
        sect_pr.append(margins)
        for side in ("top", "right", "bottom", "left"):
            margins.set(qn(f"w:{side}"), str(1440))
        margins.set(qn("w:header"), str(720))
        margins.set(qn("w:footer"), str(720))
        margins.set(qn("w:gutter"), "0")
    if landscape:
        line_number_types = sect_pr.findall(qn("w:lnNumType"))
        if len(line_number_types) != 1:
            raise RuntimeError(
                "Landscape base section must contain exactly one line-number element",
            )
        line_number_types[0].set(qn("w:distance"), "72")
        page_size.set(qn("w:w"), str(max(page_width, page_height)))
        page_size.set(qn("w:h"), str(min(page_width, page_height)))
        page_size.set(qn("w:orient"), "landscape")
        margins.set(qn("w:left"), str(round(0.45 * 1440)))
        margins.set(qn("w:right"), str(round(0.45 * 1440)))
        margins.set(qn("w:top"), str(round(0.55 * 1440)))
        margins.set(qn("w:bottom"), str(round(0.55 * 1440)))
    else:
        page_size.set(qn("w:w"), str(min(page_width, page_height)))
        page_size.set(qn("w:h"), str(max(page_width, page_height)))
        page_size.attrib.pop(qn("w:orient"), None)

    paragraph = OxmlElement("w:p")
    paragraph_properties = OxmlElement("w:pPr")
    paragraph_properties.append(OxmlElement("w:suppressLineNumbers"))
    paragraph_properties.append(sect_pr)
    paragraph.append(paragraph_properties)
    return paragraph


def wrap_landscape(document: Document, first_element, last_element, base_sect_pr):
    body = document.element.body
    before = section_break_paragraph(base_sect_pr, landscape=False)
    after = section_break_paragraph(base_sect_pr, landscape=True)
    body.insert(body.index(first_element), before)
    body.insert(body.index(last_element) + 1, after)


def manifest_map(manifest_path: Path) -> dict[str, list[Path]]:
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    result = {}
    for entry in payload:
        files = [Path(item["path"]) for item in entry["files"]]
        missing = [str(path) for path in files if not path.exists()]
        if missing:
            raise FileNotFoundError(f"Missing table PNG(s): {missing}")
        result[entry["key"]] = files
    return result


def figure_manifest_map(manifest_path: Path) -> dict[str, Path]:
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    result = {entry["key"]: Path(entry["path"]) for entry in payload}
    missing = [str(path) for path in result.values() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing figure PNG(s): {missing}")
    return result


def replace_main_table_float(
    document: Document,
    table_number: int,
    images: list[Path],
):
    body = document.element.body
    candidates = []
    for table in top_level_tables(document):
        if len(table.rows) != 1 or len(table.columns) != 1:
            continue
        text = clean_text(table.cell(0, 0).text)
        if text.startswith(f"Table {table_number}:"):
            candidates.append(table)
    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected one main Table {table_number} float, found {len(candidates)}",
        )

    wrapper = candidates[0]
    caption_candidates = [
        paragraph
        for paragraph in wrapper.cell(0, 0).paragraphs
        if clean_text(paragraph.text).startswith(f"Table {table_number}:")
    ]
    if len(caption_candidates) != 1:
        raise RuntimeError(f"Main Table {table_number} caption was not unique")
    caption = caption_candidates[0]
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.keep_with_next = True

    insertion_index = body.index(wrapper._tbl)
    caption._p.getparent().remove(caption._p)
    body.remove(wrapper._tbl)
    body.insert(insertion_index, caption._p)

    image_elements = []
    for index, image in enumerate(images):
        element = add_picture_paragraph(
            document,
            image,
            max_width_inches=10.55,
            max_height_inches=6.20,
            alt_text=f"Main Table {table_number}, part {index + 1} of {len(images)}",
            page_break_before=index > 0,
        )
        body.insert(insertion_index + 1 + index, element)
        image_elements.append(element)
    return caption._p, image_elements[-1]


def replace_main_figure_float(document: Document, figure_number: int):
    body = document.element.body
    candidates = []
    for table in top_level_tables(document):
        if len(table.rows) != 1 or len(table.columns) != 1:
            continue
        text = clean_text(table.cell(0, 0).text)
        if text.startswith(f"Figure {figure_number}:"):
            candidates.append(table)
    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected one main Figure {figure_number} float, found {len(candidates)}",
        )

    wrapper = candidates[0]
    cell_paragraphs = wrapper.cell(0, 0).paragraphs
    image_candidates = [
        paragraph
        for paragraph in cell_paragraphs
        if paragraph._p.xpath(".//w:drawing")
    ]
    caption_candidates = [
        paragraph
        for paragraph in cell_paragraphs
        if clean_text(paragraph.text).startswith(f"Figure {figure_number}:")
    ]
    if len(image_candidates) != 1 or len(caption_candidates) != 1:
        raise RuntimeError(
            f"Main Figure {figure_number} image/caption was not unique",
        )

    image = image_candidates[0]
    caption = caption_candidates[0]
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.page_break_before = True
    image.paragraph_format.keep_with_next = True
    image.paragraph_format.keep_together = True
    image.paragraph_format.space_before = Pt(0)
    image.paragraph_format.space_after = Pt(0)
    resize_drawing_paragraph(
        image,
        max_width_inches=6.25,
        max_height_inches=4.40 if figure_number == 1 else 5.35,
        alt_text=f"Main Figure {figure_number}",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_before = Pt(4)

    insertion_index = body.index(wrapper._tbl)
    image._p.getparent().remove(image._p)
    caption._p.getparent().remove(caption._p)
    body.remove(wrapper._tbl)
    body.insert(insertion_index, image._p)
    body.insert(insertion_index + 1, caption._p)


def replace_native_table(
    document: Document,
    native_table: Table,
    images: list[Path],
    label: str,
    landscape: bool,
):
    body = document.element.body
    insertion_index = body.index(native_table._tbl)
    body.remove(native_table._tbl)
    image_elements = []
    for index, image in enumerate(images):
        element = add_picture_paragraph(
            document,
            image,
            max_width_inches=10.55 if landscape else 6.25,
            max_height_inches=6.20 if landscape else 8.75,
            alt_text=f"{label}, part {index + 1} of {len(images)}",
            page_break_before=index > 0,
        )
        body.insert(insertion_index + index, element)
        image_elements.append(element)
    return image_elements


def add_do_not_compress_setting(document: Document):
    settings = document.settings._element
    if settings.find(qn("w:doNotAutoCompressPictures")) is None:
        settings.append(OxmlElement("w:doNotAutoCompressPictures"))


def suppress_linked_footer_line_numbers(document: Document):
    footer_parts = {}
    footer_references = document.element.body.xpath(
        ".//w:sectPr/w:footerReference",
    )
    for footer_reference in footer_references:
        relationship_id = footer_reference.get(qn("r:id"))
        if relationship_id is None:
            raise RuntimeError("Linked footer reference lacks a relationship ID")
        try:
            footer_part = document.part.related_parts[relationship_id]
        except KeyError as exc:
            raise RuntimeError(
                f"Could not resolve linked footer relationship {relationship_id}",
            ) from exc
        if local_name(footer_part.element) != "ftr":
            raise RuntimeError(
                f"Linked footer relationship {relationship_id} is not a footer part",
            )
        footer_parts[str(footer_part.partname)] = footer_part

    if len(footer_parts) != 2:
        raise RuntimeError(
            f"Expected exactly two linked footer parts, found {len(footer_parts)}",
        )

    for part_name, footer_part in sorted(footer_parts.items()):
        paragraphs = footer_part.element.xpath(".//w:p")
        if len(paragraphs) != 2:
            raise RuntimeError(
                f"Expected two paragraphs in {part_name}, found {len(paragraphs)}",
            )
        page_instruction_fields = [
            field
            for field in footer_part.element.xpath(".//w:instrText")
            if clean_text(field.text or "").upper() == "PAGE"
        ]
        page_simple_fields = [
            field
            for field in footer_part.element.xpath(".//w:fldSimple")
            if clean_text(field.get(qn("w:instr"), "")).upper() == "PAGE"
        ]
        if len(page_instruction_fields) + len(page_simple_fields) != 1:
            raise RuntimeError(
                f"Expected one PAGE field in {part_name}",
            )

        for paragraph in paragraphs:
            if paragraph.xpath("./w:pPr/w:suppressLineNumbers"):
                raise RuntimeError(
                    f"Footer paragraph in {part_name} already suppresses line numbers",
                )
            paragraph_properties = paragraph.find(qn("w:pPr"))
            if paragraph_properties is None:
                paragraph_properties = OxmlElement("w:pPr")
                paragraph.insert(0, paragraph_properties)
            suppression = OxmlElement("w:suppressLineNumbers")
            insertion_index = len(paragraph_properties)
            for trailing_name in ("w:rPr", "w:sectPr", "w:pPrChange"):
                trailing_element = paragraph_properties.find(qn(trailing_name))
                if trailing_element is not None:
                    insertion_index = min(
                        insertion_index,
                        paragraph_properties.index(trailing_element),
                    )
            paragraph_properties.insert(insertion_index, suppression)


def set_final_portrait_section(document: Document, base_sect_pr):
    section_properties = document.element.body.find(qn("w:sectPr"))
    if section_properties is None:
        raise RuntimeError("DOCX body has no final section properties")

    base_page_size = base_sect_pr.find(qn("w:pgSz"))
    base_margins = base_sect_pr.find(qn("w:pgMar"))
    if base_page_size is None or base_margins is None:
        raise RuntimeError("Base section lacks page size or margin properties")

    page_size = deepcopy(base_page_size)
    page_width = int(page_size.get(qn("w:w")))
    page_height = int(page_size.get(qn("w:h")))
    page_size.set(qn("w:w"), str(min(page_width, page_height)))
    page_size.set(qn("w:h"), str(max(page_width, page_height)))
    page_size.attrib.pop(qn("w:orient"), None)

    current_page_size = section_properties.find(qn("w:pgSz"))
    if current_page_size is None:
        section_properties.append(page_size)
    else:
        insertion_index = section_properties.index(current_page_size)
        section_properties.remove(current_page_size)
        section_properties.insert(insertion_index, page_size)

    margins = deepcopy(base_margins)
    current_margins = section_properties.find(qn("w:pgMar"))
    if current_margins is None:
        section_properties.append(margins)
    else:
        insertion_index = section_properties.index(current_margins)
        section_properties.remove(current_margins)
        section_properties.insert(insertion_index, margins)


def insert_supplementary_figure(
    document: Document,
    figure_number: int,
    image_path: Path,
):
    exact_number = re.compile(
        rf"^Supplementary Figure S{figure_number}(?:\.| and Table\b)",
    )
    matches = [
        paragraph
        for paragraph in top_level_paragraphs(document)
        if exact_number.match(clean_text(paragraph.text))
    ]
    if len(matches) < 2:
        raise RuntimeError(
            f"Could not resolve heading and caption for Supplementary Figure S{figure_number}",
        )
    heading = matches[0]
    caption = matches[-1]
    heading.paragraph_format.page_break_before = figure_number in {
        2,
        3,
        5,
        10,
        11,
        12,
    }
    heading.paragraph_format.keep_with_next = True
    caption.paragraph_format.keep_together = True
    if figure_number == 8:
        first_part = add_cropped_picture_paragraph(
            document,
            image_path,
            max_width_inches=6.25,
            max_height_inches=6.25,
            alt_text="Supplementary Figure S8, panels A-C",
            crop_bottom=0.57695,
        )
        continued = add_label_paragraph(
            document,
            "Supplementary Figure S8 (continued)",
            page_break_before=True,
        )
        second_part = add_cropped_picture_paragraph(
            document,
            image_path,
            max_width_inches=6.25,
            max_height_inches=6.25,
            alt_text="Supplementary Figure S8, panel D",
            crop_top=0.42305,
            keep_with_next=True,
        )
        body = document.element.body
        insertion_index = body.index(caption._p)
        body.insert(insertion_index, first_part)
        body.insert(insertion_index + 1, continued)
        body.insert(insertion_index + 2, second_part)
        return (first_part, second_part)

    height_overrides = {
        3: 4.40,
        4: 5.00,
        5: 5.00,
        6: 5.00,
        7: 5.00,
        8: 4.70,
        9: 5.50,
        10: 5.35,
        11: 5.35,
        12: 5.35,
        13: 5.25,
        14: 5.25,
        15: 5.00,
        16: 4.85,
        17: 4.85,
    }
    max_width = 4.85 if figure_number == 3 else 6.25
    max_height = height_overrides.get(figure_number, 5.75)
    image_element = add_picture_paragraph(
        document,
        image_path,
        max_width_inches=max_width,
        max_height_inches=max_height,
        alt_text=f"Supplementary Figure S{figure_number}",
        page_break_before=False,
    )
    Paragraph(image_element, document).paragraph_format.keep_with_next = True
    body = document.element.body
    body.insert(body.index(caption._p), image_element)
    return image_element


def main() -> None:
    args = parse_args()
    document = Document(args.input_docx)
    author_block = relocate_author_block(document)
    try:
        normal_style = document.styles["Normal"]
    except KeyError as exc:
        raise RuntimeError("Input DOCX lacks the Normal paragraph style") from exc
    for style_name in ("Body Text", "First Paragraph", "Compact", "Bibliography"):
        try:
            current_style = document.styles[style_name]
        except KeyError as exc:
            raise RuntimeError(f"Input DOCX lacks the {style_name} style") from exc
        while current_style is not None and current_style.style_id != normal_style.style_id:
            current_style = current_style.base_style
        if current_style is None:
            raise RuntimeError(f"{style_name} does not inherit body typography from Normal")
    for style_name, expected_size in {
        "Title": 26,
        "Heading 1": 18,
        "Heading 2": 16,
        "Heading 3": 14,
    }.items():
        try:
            actual_size = document.styles[style_name].font.size
        except KeyError as exc:
            raise RuntimeError(f"Input DOCX lacks the {style_name} style") from exc
        if actual_size is None or actual_size.pt != expected_size:
            raise RuntimeError(f"Unexpected explicit size for {style_name}")
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)
    images = manifest_map(args.png_manifest)
    figures = figure_manifest_map(args.figure_manifest)
    required_keys = {
        "main_table_1",
        "main_table_2",
        "main_table_3",
        *(f"supp_table_s{index}" for index in range(1, 11)),
        "supp_table_s11a",
        "supp_table_s11b",
        "supp_table_s12",
        "supp_table_s13",
        "supp_table_s14",
        "supp_table_s15",
    }
    if required_keys != set(images):
        missing = sorted(required_keys - set(images))
        extra = sorted(set(images) - required_keys)
        raise RuntimeError(f"PNG manifest mismatch; missing={missing}, extra={extra}")
    required_figures = {f"supp_figure_s{index}" for index in range(1, 18)}
    if required_figures != set(figures):
        missing = sorted(required_figures - set(figures))
        extra = sorted(set(figures) - required_figures)
        raise RuntimeError(f"Figure manifest mismatch; missing={missing}, extra={extra}")

    add_do_not_compress_setting(document)
    base_sect_pr = base_section_properties(document)
    landscape_ranges = []

    for number in range(1, 4):
        replace_main_figure_float(document, number)

    for number in range(1, 4):
        first, last = replace_main_table_float(
            document,
            number,
            images[f"main_table_{number}"],
        )
        landscape_ranges.append((first, last))

    supplement_keys = [
        *(f"supp_table_s{index}" for index in range(1, 11)),
        "supp_table_s11a",
        "supp_table_s11b",
        "supp_table_s12",
        "supp_table_s13",
        "supp_table_s14",
        "supp_table_s15",
    ]
    native_tables = [
        table
        for table in top_level_tables(document)
        if len(table.rows) > 1 or len(table.columns) > 1
    ]
    if len(native_tables) != len(supplement_keys):
        dimensions = [(len(table.rows), len(table.columns)) for table in native_tables]
        raise RuntimeError(
            "Expected 16 supplementary native tables, found "
            f"{len(native_tables)}: {dimensions}",
        )

    inserted = {}
    for key, native_table in zip(supplement_keys, native_tables, strict=True):
        number = key.removeprefix("supp_table_s").rstrip("ab")
        inserted[key] = replace_native_table(
            document,
            native_table,
            images[key],
            label=f"Supplementary Table S{number}",
            landscape=True,
        )

    for number in range(1, 18):
        insert_supplementary_figure(
            document,
            number,
            figures[f"supp_figure_s{number}"],
        )

    recommendation_heading = find_exact_heading(
        document,
        "Recommendation adherence",
    )
    recommendation_heading.paragraph_format.page_break_before = True
    recommendation_heading.paragraph_format.keep_with_next = True

    # Landscape groups keep each existing heading, image sequence, and caption
    # together while avoiding empty portrait sections between adjacent tables.
    for first_number, last_number in [(1, 3), (4, 4), (5, 6), (7, 8), (9, 10)]:
        first_matches = find_paragraphs(
            document,
            f"Supplementary Table S{first_number}.",
        )
        last_matches = find_paragraphs(
            document,
            f"Supplementary Table S{last_number}.",
        )
        if len(first_matches) < 2 or len(last_matches) < 2:
            raise RuntimeError(
                f"Could not resolve heading/caption range S{first_number}–S{last_number}",
            )
        first_element = first_matches[0]._p
        if first_number == 1:
            first_element = find_exact_heading(
                document,
                "Supplementary Information",
            )._p
        elif first_number == 5:
            first_element = find_exact_heading(
                document,
                "Daily architecture",
            )._p
        landscape_ranges.append((first_element, last_matches[-1]._p))

    # Combined figure-and-table sections S11–S15 retain their portrait figures.
    # Add explicit Word-only labels before the following landscape table pages.
    later_table_labels = {
        "supp_table_s11a": "Supplementary Table S11. Light-exposure behaviour and awareness",
        "supp_table_s12": "Supplementary Table S12. Visual light sensitivity",
        "supp_table_s13": "Supplementary Table S13. Chronotype and timing",
        "supp_table_s14": "Supplementary Table S14. Age and biological sex",
        "supp_table_s15": "Supplementary Table S15. Biological-sex-specific daily patterns and global tests",
    }
    body = document.element.body
    later_labels = {}
    for key, label in later_table_labels.items():
        first_image = inserted[key][0]
        label_element = add_label_paragraph(document, label, page_break_before=False)
        body.insert(body.index(first_image), label_element)
        later_labels[key] = label_element

    continued = add_label_paragraph(
        document,
        "Supplementary Table S11 (continued).",
        page_break_before=False,
    )
    body.insert(body.index(inserted["supp_table_s11b"][0]), continued)

    landscape_ranges.extend(
        [
            (
                later_labels["supp_table_s11a"],
                inserted["supp_table_s11b"][-1],
            ),
            (
                later_labels["supp_table_s12"],
                inserted["supp_table_s12"][-1],
            ),
            (
                later_labels["supp_table_s13"],
                inserted["supp_table_s13"][-1],
            ),
            (
                later_labels["supp_table_s14"],
                inserted["supp_table_s14"][-1],
            ),
            (
                later_labels["supp_table_s15"],
                inserted["supp_table_s15"][-1],
            ),
        ],
    )

    for first, last in landscape_ranges:
        wrap_landscape(document, first, last, base_sect_pr)

    set_final_portrait_section(document, base_sect_pr)
    suppress_linked_footer_line_numbers(document)
    display_bookmarks = add_missing_display_bookmarks(document)

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)

    final = Document(args.output_docx)
    remaining_native = [
        table
        for table in top_level_tables(final)
        if len(table.rows) > 1 or len(table.columns) > 1
    ]
    if remaining_native:
        raise RuntimeError("One or more lossy native gt tables remain in output")
    if top_level_tables(final):
        raise RuntimeError("One or more one-cell float wrapper tables remain in output")
    if len(final.inline_shapes) != 53:
        raise RuntimeError(
            f"Expected 53 images in final DOCX, found {len(final.inline_shapes)}",
        )
    expected_orientations = [
        orientation
        for pair in ((WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE),) * 13
        for orientation in pair
    ] + [WD_ORIENT.PORTRAIT]
    actual_orientations = [section.orientation for section in final.sections]
    if actual_orientations != expected_orientations:
        raise RuntimeError(
            "Unexpected Word section orientation sequence: "
            f"{actual_orientations}",
        )
    print(
        json.dumps(
            {
                "output": str(args.output_docx.resolve()),
                "author_block": author_block,
                "display_bookmarks": display_bookmarks,
                "sections": len(final.sections),
                "top_level_tables": len(top_level_tables(final)),
                "inline_shapes": len(final.inline_shapes),
            },
            indent=2,
        ),
    )


if __name__ == "__main__":
    main()
