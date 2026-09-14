#!/usr/bin/env python3
"""Structural and protected-content checks, not scientific verification."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from lxml import html

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts/manuscript_nature_health"))
from export_editable_tables import LABELS, compact, visible_text

AUDIT = ROOT / "audit/manuscript_nature_health/revision_2026_09_11"
MANUSCRIPT = ROOT / "manuscript/R0_NatHealth"
SOURCE = MANUSCRIPT / "ZaunerEtAl2026_NatHealth_phase3_brown.qmd"
HTML = AUDIT / "manuscript_render/ZaunerEtAl2026_NatHealth_phase3_brown.html"
DOCX = MANUSCRIPT / "_output/ZaunerEtAl2026_NatHealth_revision_2026_09_11.docx"
TABLE_DIR = MANUSCRIPT / "editable_tables"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def paragraphs(text):
    matches = list(re.finditer(r"<!-- (P-[A-Z]\d+[A-Z]?) -->", text))
    return {match.group(1): text[match.end():matches[i+1].start() if i+1<len(matches) else len(text)]
            for i,match in enumerate(matches)}


def media_hashes(path):
    with ZipFile(path) as archive:
        return sorted(hashlib.sha256(archive.read(name)).hexdigest()
                      for name in archive.namelist() if name.startswith("word/media/"))


def main():
    checks = {}
    pre = (AUDIT / "preimages" / SOURCE.name).read_text()
    current = SOURCE.read_text()
    old_map, new_map = paragraphs(pre), paragraphs(current)
    checks["paragraph_identifiers_unchanged"] = list(old_map) == list(new_map)
    changed = [key for key in old_map if old_map[key] != new_map.get(key)]
    checks["only_site_design_paragraph_changed"] = changed == ["P-M01"]
    added = "The Kumasi (GH) and Izmir (TR) collections are described in site-specific data notes [@agbeshie2025; @akgun2026]. "
    checks["site_design_change_is_only_data_note_citation"] = new_map["P-M01"].replace(added, "") == old_map["P-M01"]
    old_abstract = re.search(r"^abstract: \|\n((?:  .*\n)+)", pre, re.M)
    new_abstract = re.search(r"^# Abstract\n\n(.*?)\n\n# Introduction", current, re.M | re.S)
    checks["abstract_text_unchanged"] = compact(old_abstract.group(1)) == compact(new_abstract.group(1))
    checks["title_unchanged"] = re.search(r'^title: .*$',pre,re.M).group() == re.search(r'^title: .*$',current,re.M).group()
    checks["abstract_and_introduction_are_level_one_headings"] = bool(new_abstract)
    old_citations = set(re.findall(r"@([\w.-]+)", pre))
    new_citations = set(re.findall(r"@([\w.-]+)", current))
    checks["no_citation_key_removed"] = old_citations <= new_citations
    checks["only_two_data_note_citation_keys_added"] = new_citations - old_citations == {"agbeshie2025", "akgun2026"}
    checks["repository_citations_retained"] = all("@"+key in current for key in ("akuffo2025", "didikoglu2025"))

    before_html = html.parse(str(MANUSCRIPT / "_output" / HTML.name))
    after_html = html.parse(str(HTML))
    table_xpath = "//table[contains(concat(' ',normalize-space(@class),' '),' gt_table ')]"
    before_tables, after_tables = before_html.xpath(table_xpath), after_html.xpath(table_xpath)
    checks["all_19_html_tables_unchanged"] = (len(before_tables) == len(after_tables) == 19 and
        [compact(visible_text(table)) for table in before_tables] ==
        [compact(visible_text(table)) for table in after_tables])
    before_images = before_html.xpath("//img/@src")
    after_images = after_html.xpath("//img/@src")
    checks["all_html_image_sources_unchanged"] = before_images == after_images
    checks["html_images_embedded"] = all(source.startswith("data:") for source in after_images)
    ids = set(after_html.xpath("//*[@id]/@id"))
    anchors = after_html.xpath("//a[starts-with(@href,'#')]/@href")
    unresolved_html = sorted(set(target[1:] for target in anchors if len(target)>1) - ids)
    checks["html_internal_links_resolve"] = not unresolved_html
    for key in ("abstract", "introduction"):
        checks[f"html_{key}_heading"] = len(after_html.xpath(f"//section[@id='{key}']/h1")) == 1
    for key in ("agbeshie2025", "akgun2026"):
        checks[f"html_reference_{key}"] = bool(after_html.xpath(f"//*[@id='ref-{key}']"))

    document = Document(DOCX)
    h1s = [p for p in document.paragraphs if p.style.name == "Heading 1"]
    checks["every_word_top_level_section_starts_new_page"] = all(p.paragraph_format.page_break_before for p in h1s)
    checks["word_body_is_11_point_Arial"] = document.styles["Normal"].font.name == "Arial" and document.styles["Normal"].font.size.pt == 11
    checks["word_has_all_53_existing_display_images"] = len(document.inline_shapes) == 53
    checks["word_display_media_unchanged"] = media_hashes(DOCX) == media_hashes(ROOT / "_build/nathealth/ZaunerEtAl2026_NatHealth_phase3_brown.docx")
    root = document.element
    targets = root.xpath(".//w:hyperlink/@w:anchor")
    bookmark_names = root.xpath(".//w:bookmarkStart/@w:name")
    checks["word_internal_links_resolve"] = not (set(targets) - set(bookmark_names))
    checks["word_bookmark_names_unique"] = len(bookmark_names) == len(set(bookmark_names))
    ptexts = [p.text for p in document.paragraphs]
    abstract_index = ptexts.index("Abstract")
    checks["correspondence_precedes_abstract"] = any("Correspondence:" in text for text in ptexts[:abstract_index])
    checks["ai_statement_preserved"] = "Artificial-intelligence assistance" in ptexts
    for heading in ("Acknowledgements", "Funding", "Author contributions", "Competing interests", "Data availability", "Code availability"):
        checks[f"word_declaration_{heading}"] = heading in ptexts

    manifest = json.loads((TABLE_DIR / "table_manifest.json").read_text())
    checks["19_editable_table_documents"] = len(manifest) == len(LABELS) == 19
    for record in manifest:
        path = Path(record["path"])
        table_doc = Document(path)
        checks[record["label"]+"_one_native_table"] = len(table_doc.tables) == 1
        checks[record["label"]+"_hash_exact"] = sha(path) == record["sha256"]
        checks[record["label"]+"_source_cell_text_exact"] = record["all_displayed_cell_text_exact"]
        checks[record["label"]+"_notes_and_caption_exact"] = record["all_footnote_and_caption_text_exact"]
        checks[record["label"]+"_single_spacing"] = all(
            paragraph.paragraph_format.line_spacing == 1.0
            for row in table_doc.tables[0].rows for cell in row.cells for paragraph in cell.paragraphs)
        checks[record["label"]+"_no_row_splitting"] = all(
            bool(row._tr.xpath("./w:trPr/w:cantSplit")) for row in table_doc.tables[0].rows)

    protected = {
        "_build/nathealth/index.html": "9c1beea41a0b203b8a8b032eade2ba7e4ab5067857d39e296a8503231af46eec",
        "_build/nathealth/ZaunerEtAl2026_NatHealth_phase3_brown.docx": "74193a7a787ea18d70933b1742e588304a6bf6050c8b43e8a0a0cc352b44a2b8",
        # The local Word resave is a distinct, pre-existing 3 September file.
        # It is protected, not used as the accepted media authority.
        "manuscript/R0_NatHealth/_output/ZaunerEtAl2026_NatHealth_phase3_brown.docx": "45c866bec516919da064d7e86adbda57e6435db7b04f7ae10032f59000c77f79",
        "assets/reference.docx": "8c0cf634a4958aa05412de3f5f15aa92c40cbb5331a2319f761530472697e15f",
    }
    for path, expected in protected.items():
        checks["protected_"+path] = sha(ROOT / path) == expected
    provenance = {
        "operation": "Document conversion only; no scientific computation",
        "status": "Independent revisions complete; Brown replacement package pending",
        "source_html": str(HTML), "source_html_sha256": sha(HTML),
        "source_qmd": str(SOURCE), "source_qmd_sha256": sha(SOURCE),
        "reference_docx_sha256": sha(ROOT / "assets/reference.docx"),
        "native_tables": len(manifest),
        "brown_note": "Table 2 and associated Brown outputs retain the previous accepted content; refresh or reconfirm after the grouping amendment is accepted.",
    }
    (TABLE_DIR / "export_provenance.json").write_text(json.dumps(provenance, indent=2)+"\n")
    with ZipFile(TABLE_DIR / "editable_manuscript_tables.zip", "w", ZIP_DEFLATED) as archive:
        for record in manifest:
            archive.write(record["path"], Path(record["path"]).name)
        for name in ("README.md", "table_manifest.json", "export_provenance.json"):
            archive.write(TABLE_DIR/name, name)
    result = {
        "checks": checks,
        "all_passed": all(checks.values()),
        "changed_paragraph_ids": changed,
        "unresolved_html_links": unresolved_html,
        "word_new_page_sections": [p.text for p in h1s],
        "word_internal_hyperlinks": len(targets),
        "word_bookmarks": len(bookmark_names),
        "sha256": {str(path.relative_to(ROOT)): sha(path) for path in
                   (SOURCE, HTML, DOCX, MANUSCRIPT/"references_merged.bib", TABLE_DIR/"editable_manuscript_tables.zip")},
    }
    (AUDIT / "structural_validation.json").write_text(json.dumps(result, indent=2)+"\n")
    print(json.dumps({"passed": sum(checks.values()), "checks": len(checks),
                      "failed": [key for key,value in checks.items() if not value]}, indent=2))
    if not all(checks.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
